from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


SOURCE = Path(r"C:\Users\ZBook\Downloads\INVENTARIO RAGUILEX.xlsx")
OUTPUT_DIR = Path("outputs/raguilex_phase1")
ANALYSIS_JSON = OUTPUT_DIR / "inventory_analysis.json"
REPORT_MD = OUTPUT_DIR / "diagnostico_raguilex_fase1.md"
CLEAN_XLSX = OUTPUT_DIR / "INVENTARIO_RAGUILEX_limpio_para_importacion.xlsx"
CLEAN_CSV = OUTPUT_DIR / "inventory_clean.csv"
ISSUES_CSV = OUTPUT_DIR / "cleaning_report.csv"


CODE_RE = re.compile(r"^TN-\d{5}$", re.IGNORECASE)
PRICE_WORD_RE = re.compile(r"(bs\.?|bob|bolivianos?)", re.IGNORECASE)


def strip_accents(value: str) -> str:
    return "".join(
        ch
        for ch in unicodedata.normalize("NFKD", value)
        if not unicodedata.combining(ch)
    )


def squash_spaces(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\u00a0", " ")).strip()


def normalize_header(value: Any) -> str:
    text = squash_spaces("" if value is None else str(value))
    text = strip_accents(text).lower()
    text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")
    return text or "sin_nombre"


def normalize_text(value: Any) -> tuple[str | None, bool]:
    if value is None:
        return None, False
    text = str(value)
    cleaned = squash_spaces(text)
    if cleaned == "":
        return None, text != cleaned
    return cleaned, cleaned != text


def title_case_brand(value: Any) -> tuple[str | None, bool]:
    cleaned, changed = normalize_text(value)
    if cleaned is None:
        return None, changed
    lowered = cleaned.lower()
    known = {
        "nike": "Nike",
        "adidas": "Adidas",
        "puma": "Puma",
        "reebok": "Reebok",
        "vans": "Vans",
        "converse": "Converse",
        "fila": "Fila",
        "new balance": "New Balance",
        "under armour": "Under Armour",
        "jordan": "Jordan",
        "asics": "Asics",
    }
    if lowered in known:
        return known[lowered], True if known[lowered] != cleaned else changed
    titled = " ".join(part.capitalize() for part in lowered.split(" "))
    return titled, changed or titled != cleaned


def normalize_condition(value: Any) -> tuple[str | None, bool, bool]:
    cleaned, changed = normalize_text(value)
    if cleaned is None:
        return None, changed, False
    key = strip_accents(cleaned).lower()
    mapping = {
        "premium": "Premium",
        "excelente": "Excelente",
        "exelente": "Excelente",
        "muy bueno": "Muy bueno",
        "muy buena": "Muy bueno",
        "bueno": "Bueno",
        "buena": "Bueno",
        "regular": "Regular",
    }
    if key in mapping:
        normalized = mapping[key]
        return normalized, changed or normalized != cleaned, False
    return cleaned, changed, True


def normalize_status(value: Any) -> tuple[str | None, str | None, bool, bool]:
    cleaned, changed = normalize_text(value)
    if cleaned is None:
        return "AVAILABLE", "Disponible", changed, False
    key = strip_accents(cleaned).lower()
    mapping = {
        "disponible": ("AVAILABLE", "Disponible"),
        "available": ("AVAILABLE", "Disponible"),
        "reservado": ("RESERVED", "Reservado"),
        "reserved": ("RESERVED", "Reservado"),
        "vendido": ("SOLD", "Vendido"),
        "sold": ("SOLD", "Vendido"),
        "mantenimiento": ("MAINTENANCE", "Mantenimiento/Limpieza"),
        "limpieza": ("MAINTENANCE", "Mantenimiento/Limpieza"),
        "retirado": ("REMOVED", "Retirado"),
        "removed": ("REMOVED", "Retirado"),
    }
    if key in mapping:
        code, label = mapping[key]
        return code, label, changed or label != cleaned, False
    return None, cleaned, changed, True


def parse_decimal(value: Any) -> tuple[str | None, bool, bool]:
    if value is None or value == "":
        return None, False, False
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f"{Decimal(str(value)).quantize(Decimal('0.01'))}", False, False
    original = str(value)
    text = squash_spaces(original)
    if text == "":
        return None, original != text, False
    cleaned = PRICE_WORD_RE.sub("", text)
    if re.search(r"[A-Za-z]", cleaned):
        return None, True, True
    cleaned = cleaned.replace(" ", "")
    if "," in cleaned and "." in cleaned:
        if cleaned.rfind(",") > cleaned.rfind("."):
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:
            cleaned = cleaned.replace(",", "")
    elif "," in cleaned:
        cleaned = cleaned.replace(",", ".")
    cleaned = re.sub(r"[^0-9.\-]", "", cleaned)
    if cleaned in {"", "-", ".", "-."}:
        return None, True, True
    try:
        number = Decimal(cleaned).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None, True, True
    return f"{number}", True if str(value) != f"{number}" else False, False


@dataclass
class SheetData:
    name: str
    header_row: int | None
    headers: list[str]
    normalized_headers: list[str]
    rows: list[dict[str, Any]]
    raw_rows_count: int
    blank_columns: list[str]


def cell_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def detect_header_row(ws) -> int | None:
    best_row = None
    best_score = 0
    for row_idx in range(1, min(ws.max_row, 25) + 1):
        values = [cell.value for cell in ws[row_idx]]
        non_empty = [v for v in values if v not in (None, "")]
        strings = [v for v in non_empty if isinstance(v, str)]
        score = len(non_empty) + len(strings)
        if len(non_empty) >= 2 and score > best_score:
            best_score = score
            best_row = row_idx
    return best_row


def load_sheets() -> list[SheetData]:
    wb = load_workbook(SOURCE, data_only=False)
    sheets: list[SheetData] = []
    for ws in wb.worksheets:
        header_row = detect_header_row(ws)
        if header_row is None:
            sheets.append(SheetData(ws.title, None, [], [], [], ws.max_row, []))
            continue
        raw_headers = [cell_value(c.value) for c in ws[header_row]]
        last_col = len(raw_headers)
        while last_col and raw_headers[last_col - 1] in (None, ""):
            last_col -= 1
        raw_headers = raw_headers[:last_col]
        headers = [squash_spaces(str(v)) if v not in (None, "") else f"Columna {i+1}" for i, v in enumerate(raw_headers)]
        normalized_headers = []
        seen = Counter()
        for header in headers:
            base = normalize_header(header)
            seen[base] += 1
            normalized_headers.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
        rows: list[dict[str, Any]] = []
        for row_idx in range(header_row + 1, ws.max_row + 1):
            values = [cell_value(ws.cell(row_idx, col_idx).value) for col_idx in range(1, last_col + 1)]
            if all(v in (None, "") for v in values):
                continue
            row = {"__sheet": ws.title, "__excel_row": row_idx}
            row.update({normalized_headers[i]: values[i] for i in range(len(normalized_headers))})
            rows.append(row)
        blank_columns = []
        for idx, header in enumerate(normalized_headers):
            if all(row.get(header) in (None, "") for row in rows):
                blank_columns.append(headers[idx])
        sheets.append(SheetData(ws.title, header_row, headers, normalized_headers, rows, len(rows), blank_columns))
    return sheets


def find_column(headers: list[str], candidates: list[str]) -> str | None:
    header_set = set(headers)
    for candidate in candidates:
        if candidate in header_set:
            return candidate
    for header in headers:
        for candidate in candidates:
            if candidate in header:
                return header
    return None


def compact_counts(values: list[Any], limit: int = 30) -> dict[str, int]:
    counter = Counter()
    for value in values:
        cleaned, _ = normalize_text(value)
        if cleaned is not None:
            counter[cleaned] += 1
    return dict(counter.most_common(limit))


def analyze() -> dict[str, Any]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    sheets = load_sheets()
    data_sheets = [s for s in sheets if s.rows]
    if not data_sheets:
        raise RuntimeError("No data rows found in workbook.")
    main = max(data_sheets, key=lambda s: len(s.rows))
    headers = main.normalized_headers
    code_col = find_column(headers, ["codigo", "cod", "internal_code", "id", "sku"])
    brand_col = find_column(headers, ["marca", "brand"])
    model_col = find_column(headers, ["modelo", "model", "descripcion", "producto", "nombre"])
    size_col = find_column(headers, ["talla", "size", "numero"])
    color_col = find_column(headers, ["color"])
    cost_col = find_column(headers, ["costo", "cost", "precio_compra", "compra"])
    price_col = find_column(headers, ["precio_venta", "venta", "sale_price", "precio"])
    condition_col = find_column(headers, ["condicion", "condition", "estado_producto"])
    status_col = find_column(headers, ["estado", "status"])
    barcode_col = find_column(headers, ["codigo_original", "barcode", "upc"])
    origin_col = find_column(headers, ["origen", "procedencia"])
    entry_date_col = find_column(headers, ["fecha_ingreso", "fecha"])
    location_col = find_column(headers, ["ubicacion", "location"])
    observations_col = find_column(headers, ["descripcion_observaciones", "observaciones", "descripcion"])
    photo_col = find_column(headers, ["nombre_archivo_foto", "foto", "imagen"])

    issues: list[dict[str, Any]] = []
    clean_rows: list[dict[str, Any]] = []
    codes: list[str] = []
    duplicate_keys = defaultdict(list)
    duplicate_barcodes = defaultdict(list)
    product_like_rows = []

    for idx, row in enumerate(main.rows, start=1):
        non_code_values = [
            row.get(h)
            for h in headers
            if h != code_col and row.get(h) not in (None, "")
        ]
        is_placeholder = len(non_code_values) == 0
        if not is_placeholder:
            product_like_rows.append(row["__excel_row"])
        clean: dict[str, Any] = {
            "source_sheet": row["__sheet"],
            "source_excel_row": row["__excel_row"],
            "row_kind": "placeholder_code_only" if is_placeholder else "product_candidate",
        }
        originals = {h: row.get(h) for h in headers}
        code_original = row.get(code_col) if code_col else None
        code_clean, code_changed = normalize_text(code_original)
        if code_clean:
            code_clean = code_clean.upper()
            codes.append(code_clean)
            if not CODE_RE.match(code_clean):
                issues.append(issue(row, code_col, code_original, code_clean, "WARN", "codigo_formato_no_estandar"))
            elif code_changed or code_clean != str(code_original):
                issues.append(issue(row, code_col, code_original, code_clean, "INFO", "codigo_espacios_mayusculas"))
        else:
            issues.append(issue(row, code_col or "codigo", code_original, None, "ERROR", "codigo_faltante"))
        clean["internal_code"] = code_clean

        if brand_col:
            brand, changed = title_case_brand(row.get(brand_col))
            if changed:
                issues.append(issue(row, brand_col, row.get(brand_col), brand, "INFO", "marca_normalizada"))
            clean["brand_name"] = brand
        else:
            clean["brand_name"] = None
            issues.append(issue(row, "marca", None, None, "WARN", "columna_marca_no_detectada"))

        model_value, model_changed = normalize_text(row.get(model_col) if model_col else None)
        if model_changed:
            issues.append(issue(row, model_col, row.get(model_col), model_value, "INFO", "texto_espacios_normalizados"))
        clean["model"] = model_value

        size_value, size_changed = normalize_text(row.get(size_col) if size_col else None)
        if size_changed:
            issues.append(issue(row, size_col, row.get(size_col), size_value, "INFO", "talla_espacios_normalizados"))
        clean["size"] = size_value

        color_value, color_changed = normalize_text(row.get(color_col) if color_col else None)
        if color_changed:
            issues.append(issue(row, color_col, row.get(color_col), color_value, "INFO", "color_espacios_normalizados"))
        clean["color"] = color_value

        cost, cost_changed, cost_error = parse_decimal(row.get(cost_col) if cost_col else None)
        if cost_error:
            issues.append(issue(row, cost_col, row.get(cost_col), cost, "ERROR", "costo_no_numerico"))
        elif cost_changed:
            issues.append(issue(row, cost_col, row.get(cost_col), cost, "INFO", "costo_convertido_decimal"))
        clean["purchase_cost"] = cost

        price, price_changed, price_error = parse_decimal(row.get(price_col) if price_col else None)
        if price_error:
            issues.append(issue(row, price_col, row.get(price_col), price, "ERROR", "precio_no_numerico"))
        elif price_changed:
            issues.append(issue(row, price_col, row.get(price_col), price, "INFO", "precio_convertido_decimal"))
        clean["sale_price"] = price

        condition, condition_changed, unknown_condition = normalize_condition(row.get(condition_col) if condition_col else None)
        if unknown_condition:
            issues.append(issue(row, condition_col, row.get(condition_col), condition, "WARN", "condicion_requiere_revision"))
        elif condition_changed:
            issues.append(issue(row, condition_col, row.get(condition_col), condition, "INFO", "condicion_normalizada"))
        clean["condition_name"] = condition

        status_code, status_label, status_changed, unknown_status = normalize_status(row.get(status_col) if status_col else None)
        if unknown_status:
            issues.append(issue(row, status_col, row.get(status_col), status_label, "WARN", "estado_inventario_requiere_revision"))
        elif status_changed:
            issues.append(issue(row, status_col, row.get(status_col), status_label, "INFO", "estado_normalizado"))
        clean["inventory_status"] = status_code
        if not is_placeholder and status_code is None:
            issues.append(issue(row, status_col, row.get(status_col) if status_col else None, status_label, "ERROR", "estado_inventario_invalido"))

        barcode_value, barcode_changed = normalize_text(row.get(barcode_col) if barcode_col else None)
        if barcode_changed:
            issues.append(issue(row, barcode_col, row.get(barcode_col), barcode_value, "INFO", "codigo_original_espacios_normalizados"))
        clean["barcode"] = barcode_value
        if barcode_value:
            duplicate_barcodes[barcode_value].append(row["__excel_row"])

        origin_value, origin_changed = normalize_text(row.get(origin_col) if origin_col else None)
        if origin_changed:
            issues.append(issue(row, origin_col, row.get(origin_col), origin_value, "INFO", "origen_espacios_normalizados"))
        clean["origin"] = origin_value

        date_value = row.get(entry_date_col) if entry_date_col else None
        clean["entry_date"] = date_value.isoformat()[:10] if isinstance(date_value, datetime) else (squash_spaces(str(date_value)) if date_value not in (None, "") else None)

        location_value, location_changed = normalize_text(row.get(location_col) if location_col else None)
        if location_changed:
            issues.append(issue(row, location_col, row.get(location_col), location_value, "INFO", "ubicacion_espacios_normalizados"))
        clean["location_name"] = location_value or "Raguilex Principal"

        observations_value, observations_changed = normalize_text(row.get(observations_col) if observations_col else None)
        if observations_changed:
            issues.append(issue(row, observations_col, row.get(observations_col), observations_value, "INFO", "observaciones_espacios_normalizados"))
        clean["observations"] = observations_value

        photo_value, photo_changed = normalize_text(row.get(photo_col) if photo_col else None)
        if photo_changed:
            issues.append(issue(row, photo_col, row.get(photo_col), photo_value, "INFO", "foto_espacios_normalizados"))
        clean["image_filename"] = photo_value

        clean["category_name"] = "Tenis/Calzado"
        clean["currency"] = "BOB"
        clean["original_payload_json"] = json.dumps(originals, ensure_ascii=False, default=str)
        clean_rows.append(clean)
        duplicate_keys[code_clean].append(row["__excel_row"])

    code_dupes = {code: rows for code, rows in duplicate_keys.items() if code and len(rows) > 1}
    for code, rows in code_dupes.items():
        for excel_row in rows:
            issues.append(
                {
                    "sheet": main.name,
                    "excel_row": excel_row,
                    "field_original": code_col or "codigo",
                    "value_original": code,
                    "value_normalized": code,
                    "severity": "ERROR",
                    "action": "codigo_duplicado_en_excel",
                }
            )

    barcode_dupes = {barcode: rows for barcode, rows in duplicate_barcodes.items() if len(rows) > 1}
    for barcode, rows in barcode_dupes.items():
        for excel_row in rows:
            issues.append(
                {
                    "sheet": main.name,
                    "excel_row": excel_row,
                    "field_original": barcode_col or "codigo_original",
                    "value_original": barcode,
                    "value_normalized": barcode,
                    "severity": "WARN",
                    "action": "codigo_original_duplicado",
                }
            )

    required = ["internal_code", "brand_name", "size", "condition_name", "purchase_cost", "sale_price", "inventory_status"]
    incomplete_rows = []
    valid_rows = []
    error_rows = {i["excel_row"] for i in issues if i["severity"] == "ERROR"}
    for clean in clean_rows:
        if clean["row_kind"] == "placeholder_code_only":
            continue
        missing = [field for field in required if clean.get(field) in (None, "")]
        if missing:
            incomplete_rows.append({"excel_row": clean["source_excel_row"], "missing": missing, "code": clean.get("internal_code")})
            for field in missing:
                issues.append(
                    {
                        "sheet": clean["source_sheet"],
                        "excel_row": clean["source_excel_row"],
                        "field_original": field,
                        "value_original": None,
                        "value_normalized": None,
                        "severity": "ERROR",
                        "action": f"{field}_faltante",
                    }
                )
        if not missing and clean.get("internal_code") not in code_dupes and clean["source_excel_row"] not in error_rows:
            valid_rows.append(clean["source_excel_row"])

    analysis = {
        "source": str(SOURCE),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "sheets": [
            {
                "name": s.name,
                "header_row": s.header_row,
                "headers": s.headers,
                "normalized_headers": s.normalized_headers,
                "records": len(s.rows),
                "blank_columns": s.blank_columns,
            }
            for s in sheets
        ],
        "main_sheet": main.name,
        "detected_columns": {
            "code": code_col,
            "brand": brand_col,
            "model": model_col,
            "size": size_col,
            "color": color_col,
            "purchase_cost": cost_col,
            "sale_price": price_col,
            "condition": condition_col,
            "inventory_status": status_col,
            "barcode": barcode_col,
            "origin": origin_col,
            "entry_date": entry_date_col,
            "location": location_col,
            "observations": observations_col,
            "image_filename": photo_col,
        },
        "counts": {
            "records_found": len(clean_rows),
            "product_candidate_records": len(product_like_rows),
            "placeholder_code_only_records": len(clean_rows) - len(product_like_rows),
            "valid_records": len(valid_rows),
            "incomplete_records": len(incomplete_rows),
            "duplicate_codes": len(code_dupes),
            "duplicate_barcodes": len(barcode_dupes),
            "rows_with_errors": len({i["excel_row"] for i in issues if i["severity"] == "ERROR"}),
            "rows_with_warnings": len({i["excel_row"] for i in issues if i["severity"] == "WARN"}),
            "issues_total": len(issues),
        },
        "code": {
            "total_nonblank": len(codes),
            "unique_nonblank": len(set(codes)),
            "duplicates": code_dupes,
            "barcode_duplicates": barcode_dupes,
            "invalid_format": sorted({i["value_normalized"] for i in issues if i["action"] == "codigo_formato_no_estandar"}),
            "max_numeric": max([int(c.split("-")[1]) for c in codes if CODE_RE.match(c)] or [0]),
            "next_code": f"TN-{max([int(c.split('-')[1]) for c in codes if CODE_RE.match(c)] or [0]) + 1:05d}",
        },
        "value_profiles": {
            "brands": compact_counts([r.get(brand_col) for r in main.rows]) if brand_col else {},
            "conditions": compact_counts([r.get(condition_col) for r in main.rows]) if condition_col else {},
            "statuses": compact_counts([r.get(status_col) for r in main.rows]) if status_col else {},
            "origins": compact_counts([r.get(origin_col) for r in main.rows]) if origin_col else {},
        },
        "incomplete_rows": incomplete_rows[:200],
        "issues_sample": issues[:200],
        "output_files": {
            "clean_xlsx": str(CLEAN_XLSX.resolve()),
            "clean_csv": str(CLEAN_CSV.resolve()),
            "issues_csv": str(ISSUES_CSV.resolve()),
            "report_md": str(REPORT_MD.resolve()),
            "analysis_json": str(ANALYSIS_JSON.resolve()),
        },
    }

    pd.DataFrame(clean_rows).to_csv(CLEAN_CSV, index=False, encoding="utf-8-sig")
    pd.DataFrame(issues).to_csv(ISSUES_CSV, index=False, encoding="utf-8-sig")
    ANALYSIS_JSON.write_text(json.dumps(analysis, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    write_clean_workbook(clean_rows, issues, analysis, main)
    write_report(analysis, main)
    return analysis


def issue(row: dict[str, Any], field: str | None, original: Any, normalized: Any, severity: str, action: str) -> dict[str, Any]:
    return {
        "sheet": row["__sheet"],
        "excel_row": row["__excel_row"],
        "field_original": field or "",
        "value_original": original,
        "value_normalized": normalized,
        "severity": severity,
        "action": action,
    }


def write_clean_workbook(clean_rows: list[dict[str, Any]], issues: list[dict[str, Any]], analysis: dict[str, Any], main: SheetData) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "import_ready"
    columns = [
        "internal_code",
        "category_name",
        "brand_name",
        "model",
        "size",
        "color",
        "condition_name",
        "purchase_cost",
        "sale_price",
        "currency",
        "inventory_status",
        "barcode",
        "origin",
        "entry_date",
        "location_name",
        "observations",
        "image_filename",
        "source_sheet",
        "source_excel_row",
        "row_kind",
        "original_payload_json",
    ]
    ws.append(columns)
    for row in clean_rows:
        ws.append([row.get(col) for col in columns])
    style_header(ws)
    for col_name in ["purchase_cost", "sale_price"]:
        idx = columns.index(col_name) + 1
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row, idx)
            if cell.value not in (None, ""):
                cell.value = float(cell.value)
                cell.number_format = '#,##0.00'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    set_widths(ws, {"A": 14, "B": 18, "C": 18, "D": 28, "E": 10, "F": 18, "G": 16, "H": 14, "I": 14, "J": 10, "K": 16, "L": 22, "M": 18, "N": 14, "O": 22, "P": 34, "Q": 22, "U": 60})

    issue_ws = wb.create_sheet("cleaning_report")
    issue_cols = ["sheet", "excel_row", "field_original", "value_original", "value_normalized", "severity", "action"]
    issue_ws.append(issue_cols)
    for item in issues:
        issue_ws.append([item.get(col) for col in issue_cols])
    style_header(issue_ws)
    issue_ws.freeze_panes = "A2"
    issue_ws.auto_filter.ref = issue_ws.dimensions
    set_widths(issue_ws, {"A": 18, "B": 12, "C": 22, "D": 28, "E": 28, "F": 12, "G": 34})

    map_ws = wb.create_sheet("mapping")
    map_ws.append(["columna_excel", "campo_base_datos", "accion"])
    mapping_rows = build_mapping_rows(analysis)
    for item in mapping_rows:
        map_ws.append(item)
    style_header(map_ws)
    set_widths(map_ws, {"A": 26, "B": 28, "C": 52})

    diag_ws = wb.create_sheet("diagnostico")
    diag_rows = [
        ["archivo_fuente", analysis["source"]],
        ["hoja_principal", analysis["main_sheet"]],
        ["registros_encontrados", analysis["counts"]["records_found"]],
        ["filas_candidatas_producto", analysis["counts"]["product_candidate_records"]],
        ["filas_solo_codigo_reservado", analysis["counts"]["placeholder_code_only_records"]],
        ["registros_validos", analysis["counts"]["valid_records"]],
        ["registros_incompletos", analysis["counts"]["incomplete_records"]],
        ["codigos_duplicados", analysis["counts"]["duplicate_codes"]],
        ["codigos_originales_duplicados", analysis["counts"]["duplicate_barcodes"]],
        ["filas_con_error", analysis["counts"]["rows_with_errors"]],
        ["filas_con_advertencia", analysis["counts"]["rows_with_warnings"]],
        ["siguiente_codigo_sugerido", analysis["code"]["next_code"]],
    ]
    for row in diag_rows:
        diag_ws.append(row)
    style_header(diag_ws, rows=0)
    set_widths(diag_ws, {"A": 30, "B": 80})
    wb.save(CLEAN_XLSX)


def style_header(ws, rows: int = 1) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    white = Font(color="FFFFFF", bold=True)
    for header_row in range(1, rows + 1):
        for cell in ws[header_row]:
            cell.fill = fill
            cell.font = white


def set_widths(ws, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        if letter not in widths:
            ws.column_dimensions[letter].width = 18


def build_mapping_rows(analysis: dict[str, Any]) -> list[list[str]]:
    detected = analysis["detected_columns"]
    pairs = [
        (detected["code"], "products.internal_code", "Conservar exactamente; normalizar solo espacios/mayusculas si no altera el codigo."),
        (detected["brand"], "brands.name / products.brand_id", "Normalizar capitalizacion; crear catalogo si no existe."),
        (detected["model"], "products.model", "Limpiar espacios; no inferir modelo si falta."),
        (detected["size"], "products.size", "Guardar como texto para aceptar tallas mixtas."),
        (detected["color"], "products.color", "Limpiar espacios; opcional."),
        (detected["condition"], "conditions.name / products.condition_id", "Mapear al catalogo configurable; valores dudosos van a revision."),
        (detected["purchase_cost"], "products.purchase_cost", "Convertir a DECIMAL/NUMERIC BOB."),
        (detected["sale_price"], "products.sale_price", "Convertir a DECIMAL/NUMERIC BOB."),
        (detected["inventory_status"], "products.status", "Mapear a AVAILABLE/RESERVED/SOLD/MAINTENANCE/REMOVED."),
        (detected["barcode"], "products.barcode", "Guardar codigo original/UPC si existe; advertir duplicados."),
        (detected["origin"], "products.provenance/origin metadata", "Guardar como texto inicial; normalizar catalogo luego si se valida."),
        (detected["entry_date"], "products.created_at/import metadata", "Usar como fecha de ingreso cuando sea valida."),
        (detected["location"], "locations.name / products.location_id", "Mapear a ubicacion; usar Raguilex Principal si falta."),
        (detected["observations"], "products.observations", "Conservar observaciones descriptivas."),
        (detected["image_filename"], "product_images.image_url/import metadata", "Preparar asociacion por nombre de archivo si existe."),
        ("sin columna", "products.qr_token", "Generar en importacion con token unico opaco."),
        ("sin columna", "products.category_id", "Asignar Tenis/Calzado por defecto para esta importacion."),
        ("fila completa", "import_batches/import_rows.original_payload", "Guardar valores originales para auditoria."),
    ]
    return [[source or "no detectada", target, action] for source, target, action in pairs]


def write_report(analysis: dict[str, Any], main: SheetData) -> None:
    sheets_md = "\n".join(
        f"- {s['name']}: encabezado fila {s['header_row']}, registros {s['records']}, columnas {len(s['headers'])}"
        for s in analysis["sheets"]
    )
    columns_md = "\n".join(f"- {h}" for h in main.headers)
    profiles = analysis["value_profiles"]
    detected = analysis["detected_columns"]
    mapping_md = "\n".join(
        f"| {src or 'no detectada'} | {dst} | {act} |"
        for src, dst, act in build_mapping_rows(analysis)
    )
    counts = analysis["counts"]
    code = analysis["code"]
    model_md = """```mermaid
erDiagram
  roles ||--o{ users : assigns
  users ||--o{ products : creates
  categories ||--o{ products : groups
  brands ||--o{ products : labels
  conditions ||--o{ products : classifies
  locations ||--o{ products : stores
  products ||--o{ product_images : has
  users ||--o{ sales : sells
  locations ||--o{ sales : occurs_at
  sales ||--o{ sale_items : contains
  products ||--o{ sale_items : sold_as
  products ||--o{ inventory_movements : tracks
  users ||--o{ audit_logs : performs
```"""
    report = f"""# Diagnostico RAGUILEX Fase 1

Fuente analizada: `{analysis['source']}`

## A. Diagnostico del archivo

{sheets_md}

Hoja principal propuesta: `{analysis['main_sheet']}` por contener la mayor cantidad de registros.

## B. Columnas encontradas

{columns_md}

Columnas vacias en hoja principal: {', '.join(main.blank_columns) if main.blank_columns else 'ninguna detectada'}.

## C. Registros validos

- Registros encontrados: {counts['records_found']}
- Filas candidatas a producto: {counts['product_candidate_records']}
- Filas solo con codigo reservado: {counts['placeholder_code_only_records']}
- Registros validos para importacion directa: {counts['valid_records']}
- Registros con errores criticos: {counts['rows_with_errors']}
- Registros con advertencias: {counts['rows_with_warnings']}

## D. Registros incompletos

Registros incompletos: {counts['incomplete_records']}. Se considera incompleto si falta alguno de estos campos minimos de importacion: codigo, marca, talla, condicion, costo de compra, precio de venta o estado valido.

## E. Duplicados

Codigos internos duplicados: {counts['duplicate_codes']}. Detalle: `{json.dumps(code['duplicates'], ensure_ascii=False)}`.

Codigos originales/UPC duplicados: {counts['duplicate_barcodes']}. Detalle: `{json.dumps(code['barcode_duplicates'], ensure_ascii=False)}`.

## F. Problemas de calidad

- Total de observaciones de limpieza: {counts['issues_total']}
- Codigos con formato no estandar: {', '.join(code['invalid_format']) if code['invalid_format'] else 'ninguno detectado'}
- Siguiente codigo sugerido para altas nuevas: `{code['next_code']}`
- Valores de marca frecuentes: `{json.dumps(profiles['brands'], ensure_ascii=False)}`
- Valores de condicion frecuentes: `{json.dumps(profiles['conditions'], ensure_ascii=False)}`
- Valores de estado frecuentes: `{json.dumps(profiles['statuses'], ensure_ascii=False)}`
- Valores de origen frecuentes: `{json.dumps(profiles['origins'], ensure_ascii=False)}`

## G. Propuesta exacta de normalizacion

- Codigos: conservar `TN-00000` exactamente; solo quitar espacios y convertir a mayusculas cuando no cambie la numeracion.
- Marcas: limpiar espacios y normalizar capitalizacion (`nike`, `NIKE` -> `Nike`).
- Precios/costos: convertir texto con `Bs`, `BOB`, espacios, comas o puntos a decimal con dos posiciones; guardar como `NUMERIC`, nunca `FLOAT`.
- Condicion: mapear solo valores claros a `Premium`, `Excelente`, `Muy bueno`, `Bueno`, `Regular`; valores nuevos pasan a revision.
- Estado: separar de condicion y mapear a codigos internos estables `AVAILABLE`, `RESERVED`, `SOLD`, `MAINTENANCE`, `REMOVED`.
- Auditoria: conservar toda la fila original en `original_payload_json`.

## H. Mapeo columna Excel a base de datos

| COLUMNA EXCEL | CAMPO BASE DE DATOS | ACCION |
|---|---|---|
{mapping_md}

## I. Modelo de datos final propuesto

{model_md}

Tablas base: `users`, `roles`, `categories`, `brands`, `conditions`, `locations`, `products`, `product_images`, `sales`, `sale_items`, `inventory_movements`, `audit_logs`, `import_batches`, `import_rows`.

## J. Plan de migracion/importacion

1. Crear catalogos iniciales: roles, categorias, condiciones, ubicacion principal.
2. Crear staging de importacion con `import_batches` e `import_rows`.
3. Subir Excel, guardar hash del archivo y leer filas sin modificar fuente.
4. Normalizar en staging y mostrar vista previa con correctos, observaciones, errores y duplicados.
5. Importar solo filas sin errores criticos dentro de una transaccion.
6. Usar `internal_code` como clave idempotente; si ya existe, omitir por defecto.
7. Crear movimiento `IMPORT` por cada producto creado.
8. Generar `qr_token` unico opaco en backend, sin precio ni costo dentro del QR.
9. Mantener vendidos/retirados con estado y `active`, no borrado fisico.
"""
    REPORT_MD.write_text(report, encoding="utf-8")


if __name__ == "__main__":
    result = analyze()
    print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
